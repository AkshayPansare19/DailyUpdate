"""
Drone Data Collection API - FastAPI -> existing Excel workbook and CSV

Run:
    pip install -r requirements.txt
    uvicorn main:app --reload --port 8000

Endpoints:
    POST /upload          -> Upload an existing .xlsx workbook
    POST /message          -> Send raw message text: {"text": "...", "sender": "..."}
    GET  /records          -> Fetch all recorded logs
    GET  /download         -> Download the Excel file
    GET  /download-csv     -> Download the CSV file
    GET  /health           -> Health check
"""

import csv
from copy import copy
from datetime import date, datetime, time
from io import BytesIO
import logging
import os
import re
import threading
from typing import Optional
from zipfile import BadZipFile

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

# --------------------------------------------------------------------------
# CONFIG
# --------------------------------------------------------------------------
EXCEL_FILE = os.getenv("EXCEL_FILE", "Drone_Data.xlsx")
CSV_FILE = os.getenv("CSV_FILE", "Drone_Data.csv")
MAX_WORKBOOK_UPLOAD_BYTES = 20 * 1024 * 1024

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("drone-app")

app = FastAPI(title="Drone Data Collection API")
excel_lock = threading.RLock()
csv_lock = threading.RLock()

# --------------------------------------------------------------------------
# DATA LAYOUT
# --------------------------------------------------------------------------
COMMON = ["Sr No", "Date", "Time", "Drone Pilot Name", "District", "Substation Name", "Feeder Name", "Lineman Name"]
META = ["Sender Mobile"]

SHEETS = {
    "Starting": COMMON + META,
    "Ending": COMMON + ["Total Assets Captured"] + META,
    "Issues": COMMON + ["Issue Type", "Issue Description"] + META,
    "Daily Routine": ["Sr No", "Date", "Time", "District", "Substation Name", "Completed", "Remaining", "Description"] + META,
}

CSV_HEADERS = [
    "Sr No",
    "Type",
    "Date",
    "Time",
    "Drone Pilot Name",
    "District",
    "Substation Name",
    "Feeder Name",
    "Lineman Name",
    "Total Assets Captured",
    "Issue Type",
    "Issue Description",
    "Completed",
    "Remaining",
    "Description",
    "Sender Mobile",
]

def _format_time(raw_time: Optional[str]) -> str:
    """Format time string cleanly like '01 : 36 PM' or default to current time in that format."""
    if raw_time:
        t = _clean(raw_time)
        m = re.search(r"(\d{1,2})\s*[:.]\s*(\d{2})(?:\s*[:.]\s*\d{2})?\s*([AaPp][Mm])?", t)
        if m:
            hh, mm, ampm = m.groups()
            hh = f"{int(hh):02d}"
            ampm = f" {ampm.upper()}" if ampm else ""
            return f"{hh} : {mm}{ampm}"
        return t
    return datetime.now().strftime("%I : %M %p")


def _safe_save_workbook(wb):
    """Safely save workbook using atomic replace to prevent file corruption if locked."""
    tmp_file = f"{EXCEL_FILE}.tmp"
    wb.save(tmp_file)
    try:
        os.replace(tmp_file, EXCEL_FILE)
    except (PermissionError, OSError) as e:
        if os.path.exists(tmp_file):
            try:
                os.remove(tmp_file)
            except OSError:
                pass
        raise e


class WorkbookImportError(Exception):
    pass


def append_excel(sheet_name: str, values: dict) -> int:
    """Append a record to an existing worksheet without creating workbook structure."""
    with excel_lock:
        if not os.path.isfile(EXCEL_FILE):
            raise WorkbookImportError("Upload an existing Excel workbook before importing a message.")

        headers = SHEETS[sheet_name]
        try:
            wb = load_workbook(EXCEL_FILE, keep_vba=EXCEL_FILE.lower().endswith(".xlsm"))
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise WorkbookImportError(f"Could not open the uploaded workbook: {exc}") from exc

        try:
            existing_sheet = next(
                (name for name in wb.sheetnames if name.casefold() == sheet_name.casefold()),
                None,
            )
            if existing_sheet is None:
                raise WorkbookImportError(
                    f"The uploaded workbook does not contain an existing '{sheet_name}' sheet. "
                    "No sheet was created."
                )
            ws = wb[existing_sheet]
            columns = {}
            required_headers = {header.casefold() for header in headers}
            for column in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=column).value
                if header is not None:
                    key = str(header).strip().casefold()
                    if key in required_headers and key in columns:
                        raise WorkbookImportError(
                            f"The '{sheet_name}' sheet has duplicate column headers: {header}."
                        )
                    columns.setdefault(key, column)

            missing = [header for header in headers if header.casefold() not in columns]
            if missing:
                raise WorkbookImportError(
                    f"The '{sheet_name}' sheet is missing required column headers: {', '.join(missing)}. "
                    "No columns were added."
                )

            last_data_row = 1
            for row_no in range(2, ws.max_row + 1):
                if any(
                    ws.cell(row=row_no, column=column).value not in (None, "")
                    for column in range(1, ws.max_column + 1)
                ):
                    last_data_row = row_no
            row_no = last_data_row + 1

            serial_column = columns["sr no"]
            highest_serial = 0
            for existing_row in range(2, last_data_row + 1):
                serial = ws.cell(row=existing_row, column=serial_column).value
                if isinstance(serial, (int, float)) and not isinstance(serial, bool):
                    highest_serial = max(highest_serial, int(serial))
                elif isinstance(serial, str) and serial.strip().isdigit():
                    highest_serial = max(highest_serial, int(serial.strip()))

            values = {**values, "Sr No": highest_serial + 1}
            for header in headers:
                column = columns[header.casefold()]
                cell = ws.cell(row=row_no, column=column, value=values.get(header))
                if last_data_row >= 2:
                    previous_cell = ws.cell(row=last_data_row, column=column)
                    if previous_cell.has_style:
                        cell._style = copy(previous_cell._style)

            _safe_save_workbook(wb)
            return row_no
        except PermissionError as exc:
            raise WorkbookImportError(
                f"{EXCEL_FILE} is open or locked. Close it in Excel and try again."
            ) from exc
        finally:
            wb.close()


def append_csv(sheet_name: str, values: dict) -> int:
    """Append one record to the single CSV_FILE."""
    with csv_lock:
        file_exists = os.path.exists(CSV_FILE)
        row_count = 0
        if file_exists:
            try:
                with open(CSV_FILE, "r", encoding="utf-8-sig") as f:
                    row_count = sum(1 for line in f if line.strip())
            except Exception:
                row_count = 0

        sr_no = row_count if row_count > 0 else 1

        date_val = values.get("Date")
        if isinstance(date_val, datetime):
            date_str = date_val.strftime("%d-%m-%Y")
        else:
            date_str = str(date_val or "")

        time_str = str(values.get("Time", "") or "")

        row_data = [
            sr_no,
            sheet_name,
            date_str,
            time_str,
            values.get("Drone Pilot Name", "") or "",
            values.get("District", "") or "",
            values.get("Substation Name", "") or "",
            values.get("Feeder Name", "") or "",
            values.get("Lineman Name", "") or "",
            values.get("Total Assets Captured", "") if values.get("Total Assets Captured") is not None else "",
            values.get("Issue Type", "") or "",
            values.get("Issue Description", "") or "",
            values.get("Completed", "") if values.get("Completed") is not None else "",
            values.get("Remaining", "") if values.get("Remaining") is not None else "",
            values.get("Description", "") or "",
            values.get("Sender Mobile", "") or "",
        ]

        with open(CSV_FILE, "a", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            if not file_exists or row_count == 0:
                writer.writerow(CSV_HEADERS)
                row_data[0] = 1
            writer.writerow(row_data)

        return sr_no


# --------------------------------------------------------------------------
# MESSAGE PARSER
# --------------------------------------------------------------------------
KNOWN_KEYS = {
    "date": "date",
    "time": "time",
    "drone pilot name": "pilot",
    "district": "district",
    "substation name": "substation",
    "feeder name": "feeder",
    "lineman name": "lineman",
    "total number of assets data captured today": "total_assets",
    "issue type": "issue_type",
    "issue description": "issue_description",
    "completed": "completed",
    "remaining": "remaining",
    "description": "description",
}

DATE_FORMATS = [
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y",
    "%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d-%b-%Y", "%d-%B-%Y", "%d %b %y",
]


def _clean(v: Optional[str]) -> str:
    """Remove WhatsApp bold/italic markers, blank-line underscores, extra spaces."""
    if not v:
        return ""
    v = v.replace("*", "").replace("~", "")
    v = re.sub(r"_{2,}", "", v)  # "____" blanks
    return re.sub(r"[ \t]+", " ", v).strip()


def _parse_date(v: str):
    v = _clean(v)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(v, fmt)
        except ValueError:
            continue
    return v  # leave as text if format is unknown


def _parse_int(v: str):
    m = re.search(r"\d+", _clean(v))
    return int(m.group()) if m else _clean(v)


def extract_fields(text: str) -> dict:
    """Line-based 'Label: value' parser. Values can continue on following lines."""
    fields: dict[str, str] = {}
    current = None
    for raw in text.replace("\r", "").split("\n"):
        line = _clean(raw)
        if not line:
            continue
        if ":" in line:
            label, _, rest = line.partition(":")
            label_norm = re.sub(r"[^a-z ]", "", label.lower()).strip()
            label_norm = re.sub(r"\s+", " ", label_norm)
            if label_norm in KNOWN_KEYS:
                current = KNOWN_KEYS[label_norm]
                fields[current] = rest.strip()
                continue
            # Header-only section lines with colon but no value (e.g., 'Data Collection Progress:')
            if not rest.strip():
                current = None
                continue
        if current:
            # If current field is 'remaining' and additional lines follow, treat them as 'description'
            if current == "remaining":
                if "description" not in fields:
                    fields["description"] = line
                    current = "description"
                else:
                    fields["description"] = (fields["description"] + "\n" + line).strip()
            else:
                fields[current] = (fields[current] + "\n" + line).strip()
    return fields


def detect_type(text: str, fields: dict) -> Optional[str]:
    t = text.lower()
    # Check Daily routine / ending report first so "ending report" isn't confused with ending message
    if "ending report" in t or "daily routine" in t:
        return "Daily Routine"
    if "completed" in fields and "remaining" in fields:
        return "Daily Routine"
    if "ending message" in t:
        return "Ending"
    if "starting message" in t:
        return "Starting"
    if "issue report" in t:
        return "Issues"
    if "total_assets" in fields:
        return "Ending"
    if "issue_type" in fields:
        return "Issues"
    if "pilot" in fields:
        return "Starting"
    return None


def normalise_issue_type(raw: str) -> str:
    """Map whatever the pilot left in 'Issue Type' to a clean standard label(s)."""
    t = raw.lower()
    found = []
    if "nature" in t or "weather" in t:
        found.append("Nature/Weather Problem")
    if "not available" in t:
        found.append("Lineman Not Available")
    if "arrived late" in t or "late" in t:
        found.append("Lineman Arrived Late")
    if "not receiving" in t or "calls" in t:
        found.append("Lineman Not Receiving Calls")
    if "communication" in t or "coordination" in t:
        found.append("Communication/Coordination Issue")
    m = re.search(r"other\s*:?\s*(.*)", raw, flags=re.I)
    if m:
        detail = re.sub(r"[^\w\s/,.\-()]", "", m.group(1)).strip()
        found.append(f"Other: {detail}" if detail else "Other")
    return " | ".join(dict.fromkeys(found)) if found else _clean(raw)


def parse_message(text: str) -> Optional[tuple[str, dict]]:
    """Returns (sheet_name, row_values) or None if message isn't one of our formats."""
    fields = extract_fields(text)
    kind = detect_type(text, fields)
    if not kind:
        return None

    if kind == "Daily Routine":
        row = {
            "Date": _parse_date(fields.get("date", "")),
            "Time": _format_time(fields.get("time")),
            "District": _clean(fields.get("district")),
            "Substation Name": _clean(fields.get("substation")),
            "Completed": _parse_int(fields.get("completed", "")) if fields.get("completed") else _clean(fields.get("completed")),
            "Remaining": _parse_int(fields.get("remaining", "")) if fields.get("remaining") else _clean(fields.get("remaining")),
            "Description": _clean(fields.get("description")),
        }
        return kind, row

    row = {
        "Date": _parse_date(fields.get("date", "")),
        "Time": _format_time(fields.get("time")),
        "Drone Pilot Name": _clean(fields.get("pilot")),
        "District": _clean(fields.get("district")),
        "Substation Name": _clean(fields.get("substation")),
        "Feeder Name": _clean(fields.get("feeder")),
        "Lineman Name": _clean(fields.get("lineman")),
    }
    if kind == "Ending":
        row["Total Assets Captured"] = _parse_int(fields.get("total_assets", ""))
    if kind == "Issues":
        row["Issue Type"] = normalise_issue_type(fields.get("issue_type", ""))
        row["Issue Description"] = _clean(fields.get("issue_description"))
    return kind, row


def process_text(text: str, sender: str = "", time_override: str = "") -> dict:
    parsed = parse_message(text)
    if not parsed:
        return {"saved": False, "reason": "Message format not recognised"}
    sheet, row = parsed
    if time_override.strip():
        row["Time"] = _format_time(time_override)
    row["Sender Mobile"] = sender

    excel_row = append_excel(sheet, row)
    csv_row = None
    csv_warning = None
    try:
        csv_row = append_csv(sheet, row)
    except (OSError, csv.Error, UnicodeError) as exc:
        log.exception("Excel was updated, but the CSV could not be updated.")
        csv_warning = f"The Excel workbook was updated, but the CSV file could not be updated: {exc}"

    log.info("Saved %s record: CSV row %s, Excel row %s (sender %s)", sheet, csv_row, excel_row, sender)
    res = {
        "saved": True,
        "sheet": sheet,
        "excel_row": excel_row,
        "csv_row": csv_row,
        "data": _jsonable(row),
    }
    if csv_warning:
        res["warning"] = csv_warning
    return res


def _jsonable(d: dict) -> dict:
    return {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in d.items()}


# --------------------------------------------------------------------------
# ROUTES
# --------------------------------------------------------------------------
class SimpleMessage(BaseModel):
    text: str
    sender: str = ""
    time: str = ""


@app.get("/")
def root():
    index_path = os.path.join(os.path.dirname(__file__), "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path, media_type="text/html")
    return {
        "message": "Drone Data Collection API is running",
        "docs": "/docs",
        "health": "/health",
        "download_excel": "/download",
        "download_csv": "/download-csv",
    }


@app.get("/records")
def get_records():
    """Return all saved records directly from Drone_Data.xlsx for the frontend table."""
    if not os.path.exists(EXCEL_FILE):
        return {"records": []}
    with excel_lock:
        try:
            wb = load_workbook(EXCEL_FILE, data_only=True)
            records = []
            for sheet_name in ["Starting", "Ending", "Issues", "Daily Routine"]:
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
                for r in range(2, ws.max_row + 1):
                    row_vals = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
                    if not any(row_vals.values()):
                        continue

                    date_v = row_vals.get("Date")
                    if isinstance(date_v, datetime):
                        date_str = date_v.strftime("%d/%m/%Y")
                    else:
                        date_str = str(date_v or "")

                    records.append({
                        "Type": sheet_name,
                        "Sr No": row_vals.get("Sr No", ""),
                        "Date": date_str,
                        "Time": str(row_vals.get("Time", "") or ""),
                        "Drone Pilot Name": row_vals.get("Drone Pilot Name", "") or "",
                        "District": row_vals.get("District", "") or "",
                        "Substation Name": row_vals.get("Substation Name", "") or "",
                        "Feeder Name": row_vals.get("Feeder Name", "") or "",
                        "Lineman Name": row_vals.get("Lineman Name", "") or "",
                        "Total Assets Captured": row_vals.get("Total Assets Captured", "") or "",
                        "Issue Type": row_vals.get("Issue Type", "") or "",
                        "Issue Description": row_vals.get("Issue Description", "") or "",
                        "Completed": row_vals.get("Completed", "") or "",
                        "Remaining": row_vals.get("Remaining", "") or "",
                        "Description": row_vals.get("Description", "") or "",
                        "Sender Mobile": row_vals.get("Sender Mobile", "") or "",
                    })
            return {"records": records}
        except Exception as e:
            log.exception("Error reading records from %s", EXCEL_FILE)
            return {"records": [], "error": str(e)}


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/workbook/status")
def workbook_status():
    return {
        "uploaded": os.path.isfile(EXCEL_FILE),
        "filename": os.path.basename(EXCEL_FILE),
    }


@app.post("/upload")
async def upload_workbook(request: Request, filename: str = "workbook.xlsx"):
    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an Excel .xlsx file.")

    content_length = request.headers.get("content-length")
    if content_length and content_length.isdigit() and int(content_length) > MAX_WORKBOOK_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The workbook must be 20 MB or smaller.")

    contents = await request.body()
    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded workbook is empty.")
    if len(contents) > MAX_WORKBOOK_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The workbook must be 20 MB or smaller.")

    workbook = None
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True)
        if not workbook.sheetnames:
            raise HTTPException(status_code=400, detail="The workbook does not contain any existing sheets.")
    except HTTPException:
        raise
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=f"Could not read the Excel workbook: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()

    temporary_file = f"{EXCEL_FILE}.upload.tmp"
    try:
        with excel_lock:
            with open(temporary_file, "wb") as uploaded_file:
                uploaded_file.write(contents)
            os.replace(temporary_file, EXCEL_FILE)
    except OSError as exc:
        if os.path.exists(temporary_file):
            os.remove(temporary_file)
        raise HTTPException(
            status_code=409,
            detail="Could not replace the current workbook. Close it in Excel and try again.",
        ) from exc

    return {"uploaded": True, "filename": os.path.basename(filename.replace("\\", "/"))}


@app.get("/workbook/preview")
def preview_workbook(sheet: Optional[str] = None, offset: int = 0, limit: int = 100):
    if offset < 0 or limit < 1 or limit > 200:
        raise HTTPException(status_code=400, detail="Preview offset must be non-negative and limit must be 1-200.")
    if not os.path.isfile(EXCEL_FILE):
        raise HTTPException(status_code=404, detail="Upload an Excel workbook before previewing it.")

    with excel_lock:
        workbook = None
        try:
            workbook = load_workbook(EXCEL_FILE, read_only=True, data_only=False)
            sheet_names = workbook.sheetnames
            if not sheet_names:
                raise HTTPException(status_code=400, detail="The workbook does not contain any sheets.")
            selected_sheet = next(
                (name for name in sheet_names if sheet and name.casefold() == sheet.casefold()),
                sheet_names[0] if sheet is None else None,
            )
            if selected_sheet is None:
                raise HTTPException(status_code=404, detail=f"Sheet '{sheet}' was not found.")

            worksheet = workbook[selected_sheet]
            column_count = worksheet.max_column
            headers = [_preview_value(value) for value in next(
                worksheet.iter_rows(min_row=1, max_row=1, max_col=column_count, values_only=True)
            )]
            first_row = 2 + offset
            final_row = min(worksheet.max_row, first_row + limit - 1)
            rows = []
            for row_number, row_values in enumerate(
                worksheet.iter_rows(
                    min_row=first_row,
                    max_row=final_row,
                    max_col=column_count,
                    values_only=True,
                ),
                start=first_row,
            ):
                values = [_preview_value(value) for value in row_values]
                if any(value not in (None, "") for value in values):
                    rows.append({"row_number": row_number, "values": values})
            result = {
                "filename": os.path.basename(EXCEL_FILE),
                "sheets": sheet_names,
                "sheet": selected_sheet,
                "headers": headers,
                "rows": rows,
                "offset": offset,
                "limit": limit,
                "total_rows": worksheet.max_row,
                "has_more": final_row < worksheet.max_row,
            }
            return result
        except HTTPException:
            raise
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            log.exception("Could not preview workbook %s", EXCEL_FILE)
            raise HTTPException(status_code=400, detail=f"Could not read the Excel workbook: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()


def _preview_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


@app.post("/message")
def receive_simple_message(body: SimpleMessage):
    """Process raw message text submitted via UI, script, or POST API."""
    try:
        return process_text(body.text, sender=body.sender, time_override=body.time)
    except WorkbookImportError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/download")
def download_excel():
    if not os.path.exists(EXCEL_FILE):
        raise HTTPException(404, "No data yet")
    return FileResponse(
        EXCEL_FILE,
        filename=os.path.basename(EXCEL_FILE),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )


@app.get("/download-csv")
def download_csv():
    if not os.path.exists(CSV_FILE):
        raise HTTPException(404, "No data yet")
    return FileResponse(
        CSV_FILE,
        filename=os.path.basename(CSV_FILE),
        media_type="text/csv",
        headers={
            "Cache-Control": "no-cache, no-store, must-revalidate",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )